from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "documents"
OUTPUT.mkdir(parents=True, exist_ok=True)

INK = RGBColor(0, 0, 0)
MUTED = RGBColor(45, 45, 45)
ACCENT = RGBColor(0, 0, 0)
FONT = "Arial"


def set_font(run, size=None, bold=None, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def configure_document(doc, compact=False):
    doc.core_properties.author = "Conold Thando Chisinahama"
    doc.core_properties.title = "Conold Chisinahama Resume"
    doc.core_properties.subject = "Full-stack software and product engineering"
    doc.core_properties.keywords = "TypeScript, React, Next.js, Node.js, PostgreSQL"
    doc.core_properties.comments = ""
    # Use a Unicode bullet in the document font for reliable PDF text extraction.
    for level in doc.part.numbering_part.element.findall(".//" + qn("w:lvl")):
        number_format = level.find(qn("w:numFmt"))
        if number_format is not None and number_format.get(qn("w:val")) == "bullet":
            level.find(qn("w:lvlText")).set(qn("w:val"), "\u2022")
            properties = level.find(qn("w:rPr"))
            if properties is None:
                properties = OxmlElement("w:rPr")
                level.append(properties)
            fonts = properties.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                properties.append(fonts)
            fonts.set(qn("w:ascii"), FONT)
            fonts.set(qn("w:hAnsi"), FONT)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    margin = 0.58 if compact else 0.68
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10 if compact else 10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3 if compact else 4)
    normal.paragraph_format.line_spacing = 1.0 if compact else 1.05

    for name, size, before, after in [
        ("Title", 25 if compact else 27, 0, 2),
        ("Heading 1", 11, 9 if compact else 11, 4),
        ("Heading 2", 10.5, 5, 1),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_properties = doc.styles["Title"]._element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)
    doc.styles["Heading 1"].font.all_caps = True
    doc.styles["Heading 1"].font.letter_spacing = Pt(0.7)

    if "Resume Meta" not in doc.styles:
        meta = doc.styles.add_style("Resume Meta", WD_STYLE_TYPE.PARAGRAPH)
        meta.base_style = doc.styles["Normal"]
        meta.font.name = FONT
        meta.font.size = Pt(9)
        meta.font.color.rgb = MUTED
        meta.paragraph_format.space_after = Pt(1)


def add_header(doc, subtitle):
    title = doc.add_paragraph(style="Title")
    title.add_run("Conold Thando Chisinahama")
    title_properties = title._p.get_or_add_pPr()
    paragraph_borders = title_properties.find(qn("w:pBdr"))
    if paragraph_borders is not None:
        title_properties.remove(paragraph_borders)
    role = doc.add_paragraph()
    role.paragraph_format.space_after = Pt(4)
    run = role.add_run(subtitle)
    set_font(run, 11, True, ACCENT)
    contact = doc.add_paragraph(style="Resume Meta")
    contact.add_run("Johannesburg, South Africa | +27 82 886 8556 | cthandoc@gmail.com")
    links = doc.add_paragraph(style="Resume Meta")
    for index, (label, url) in enumerate([
        ("conold-portfolio.vercel.app", "https://conold-portfolio.vercel.app"),
        ("linkedin.com/in/conold", "https://www.linkedin.com/in/conold/"),
        ("github.com/thvndx", "https://github.com/thvndx"),
    ]):
        if index:
            links.add_run(" | ")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        hyperlink.append(run)
        links._p.append(hyperlink)


def add_section(doc, title):
    doc.add_heading(title, level=1)


def add_role(doc, title, company, dates, location=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_font(run, 10.5, True)
    run = p.add_run(f" | {company}")
    set_font(run, 10.5, True, MUTED)
    meta = doc.add_paragraph(style="Resume Meta")
    text = dates if not location else f"{dates} | {location}"
    meta.add_run(text)


def add_bullet(doc, text, compact=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(1.4 if compact else 2.4)
    p.paragraph_format.line_spacing = 1.0 if compact else 1.04
    set_font(p.add_run(text), 10 if compact else 10.25)


def add_skills(doc, lines, compact=False):
    for label, values in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1 if compact else 2)
        set_font(p.add_run(f"{label}: "), 10, True)
        set_font(p.add_run(values), 10, False, MUTED)


def add_footer(doc, label):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(label)
    set_font(run, 8, False, MUTED)
    set_font(footer.add_run(" | "), 8, False, MUTED)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)


def build_concise():
    doc = Document()
    configure_document(doc, compact=True)
    add_header(doc, "Full-Stack Software Developer | Customer Success Lead")

    add_section(doc, "Profile")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run("Full-stack software developer building TypeScript applications, API integrations and analytics tools used by a four-person support operation and adjacent departments. Owns product definition, implementation, testing and deployment, drawing on customer-success leadership to solve operational problems."), 10)

    add_section(doc, "Technical Skills")
    add_skills(doc, [
        ("Development", "TypeScript, JavaScript, React, Next.js, Node.js, Python, PostgreSQL"),
        ("Delivery", "REST APIs, Supabase, Vitest, Playwright, GitHub Actions, Railway, Vercel"),
    ], compact=True)

    add_section(doc, "Experience")
    add_role(doc, "Customer Success Lead", "Connectivity and e-commerce business", "January 2023 - Present", "Remote / Johannesburg")
    for text in [
        "Built and deployed full-stack tools for customer context, quality review and performance reporting, adopted across the full support team and adjacent departments.",
        "Integrated support, call and commerce data into a customer workspace, reducing context gathering from approximately 20 minutes to 30 seconds per interaction.",
        "Automated KPI report generation and QA selection/compilation, reducing each workflow from approximately two hours to about one minute while retaining human review.",
        "Implemented explicit identity review, independent access controls and data-freshness states to prevent ambiguous or incomplete data from driving unsafe decisions.",
        "Lead two direct reports within a four-person operation; maintain measured email response and callback times below two hours and customer satisfaction above 80%.",
    ]:
        add_bullet(doc, text, compact=True)

    add_section(doc, "Selected Product Work")
    products = [
        ("Support Quality Platform", "TypeScript, React, Node.js, PostgreSQL, OpenAI. Built evidence-backed conversation scoring, sampling and manager reporting with separate manual and automated permissions."),
        ("Customer Operations Workbench", "TypeScript, React, Node.js, PostgreSQL, Redis. Integrated Front, Shopify and Aircall with bounded caching and explicit review of conflicting customer matches."),
        ("Operations Performance Hub", "TypeScript, React, PostgreSQL. Separated raw activity from KPI eligibility and added persisted, budget-aware provider synchronization."),
        ("Recalibrate", "React Native, Expo. Built a local-first planning MVP with adaptive scheduling and supervised guidance."),
    ]
    for name, description in products:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        set_font(p.add_run(f"{name}: "), 10, True)
        set_font(p.add_run(description), 10, False, MUTED)

    add_section(doc, "Earlier Experience and Education")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("Technical Sales and Support, 2020-2023 | Technician, 2019 | Technical Sales Consultant, 2015-2017"), 9.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("National Diploma in Business Information Technology, University of Johannesburg, 2014-2018"), 9.5, True)
    p = doc.add_paragraph()
    set_font(p.add_run("Peplink Certified Engineer, awarded November 2025, valid through November 2027"), 9.5)
    doc.save(OUTPUT / "conold-chisinahama-resume-concise.docx")


def build_master():
    doc = Document()
    configure_document(doc, compact=False)
    add_header(doc, "Full-Stack Software Developer | Customer Success Lead")

    add_section(doc, "Professional Summary")
    doc.add_paragraph("Full-stack software developer building operational applications, API integrations and analytics products with TypeScript, React, Node.js and PostgreSQL. Owns delivery from product definition through deployment and release verification. Combines hands-on engineering with customer-success leadership; internal tools are used across a four-person support operation and adjacent departments.")

    add_section(doc, "Core Skills")
    add_skills(doc, [
        ("Languages and frameworks", "TypeScript, JavaScript, React, Next.js, Node.js, Python, React Native, Expo"),
        ("Data and integrations", "PostgreSQL, Supabase, REST APIs, Shopify, Front, Aircall, OAuth, background jobs"),
        ("Quality and delivery", "Vitest, Playwright, GitHub Actions, Railway, Vercel, responsive and accessibility testing"),
        ("Leadership and product", "Customer success, team leadership, process design, KPI definition, vendor negotiation, product discovery"),
    ])

    add_section(doc, "Professional Experience")
    add_role(doc, "Customer Success Lead", "Connectivity and e-commerce business", "January 2023 - Present", "Remote / Johannesburg")
    for text in [
        "Built and deployed internal applications for customer context, quality review, experience analysis and performance reporting, adopted across the full support operation and adjacent departments.",
        "Integrated support, call and commerce data into a customer workspace, reducing context gathering from approximately 20 minutes to 30 seconds per interaction.",
        "Normalized provider data and persisted reporting snapshots, reducing recurring KPI report generation from approximately two hours to about one minute.",
        "Automated QA selection and report compilation, reducing the workflow from approximately two hours to about one minute while preserving human review and independent permission controls.",
        "Built evidence-backed experience scoring, reducing customer-satisfaction investigation and calculation from approximately 30 minutes to about one minute.",
        "Implemented explicit identity-conflict review, data-freshness states and provider request budgets to make uncertain data and integration limits visible.",
        "Own product definition, architecture, implementation, testing, deployment and post-release verification for the internal products delivered.",
        "Lead two direct reports within a four-person operation; maintain measured email response and callback times below two hours and customer satisfaction above 80%.",
        "Reduced recurring software costs through vendor negotiation and provider changes.",
    ]:
        add_bullet(doc, text)

    doc.add_page_break()
    add_section(doc, "Selected Product Engineering")
    product_lines = [
        ("Support Quality Platform", "TypeScript, React, Node.js, PostgreSQL, OpenAI. Built conversation scoring, automated sampling and reporting with independent manual-review and automation controls."),
        ("Customer Operations Workbench", "TypeScript, React, Node.js, PostgreSQL, Redis. Integrated Front, Shopify and Aircall; added bounded caching and explicit review of conflicting customer identities."),
        ("Operations Performance Hub", "TypeScript, React, PostgreSQL. Separated raw provider activity from KPI attribution and implemented persisted, budget-aware synchronization."),
        ("PromoStudio", "Built authentication, persisted campaign briefs, AI generation, scheduling and controlled external publishing around an approval-first workflow."),
        ("Recalibrate", "React Native, Expo. Designed and built a local-first mobile planning MVP with adaptive scheduling and supervised guidance."),
    ]
    for name, detail in product_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2.4)
        set_font(p.add_run(f"{name}: "), 9.8, True)
        set_font(p.add_run(detail), 9.8, False, MUTED)

    add_section(doc, "Earlier Experience")
    add_role(doc, "Sales Assistant - Technical Sales and Support", "Connectivity business", "May 2020 - January 2023", "Remote / Irvine, California")
    add_bullet(doc, "Provided remote technical and sales support for customers selecting, installing and troubleshooting connectivity equipment.")
    add_bullet(doc, "Translated technical constraints into clear recommendations for home and commercial environments.")

    add_role(doc, "Technician", "JJ Communications PTY LTD", "February 2019 - November 2019", "Johannesburg, South Africa")
    add_bullet(doc, "Installed, maintained and programmed nurse-call, smoke detection, access-control, public-address and evacuation systems for hospital environments across Southern Africa.")

    add_role(doc, "Technical Sales Consultant", "iStore - Apple Premium Reseller", "September 2015 - January 2017", "Johannesburg, South Africa")
    add_bullet(doc, "Helped customers choose, configure and troubleshoot technology by combining technical diagnosis with practical sales guidance.")

    add_section(doc, "Earlier Software Work")
    add_bullet(doc, "Co-built an early Python returns-rate dashboard with a colleague over approximately two to three months; later made a bounded contribution to the production platform.")
    add_bullet(doc, "Built SlykPark as a 2017 university capstone: a wallet-funded parking application connecting number-plate recognition, vehicle entry, time on site and automated payment.")

    add_section(doc, "Education")
    add_role(doc, "National Diploma - Business Information Technology", "University of Johannesburg", "2014 - 2018", "Johannesburg, South Africa")

    add_section(doc, "Certification")
    add_role(doc, "Peplink Certified Engineer", "Peplink", "Awarded 17 November 2025 | Valid through 17 November 2027")

    add_section(doc, "Languages")
    doc.add_paragraph("English - Native or bilingual | isiZulu - Professional working | Afrikaans - Limited working | French - Elementary")

    add_footer(doc, "Conold Thando Chisinahama | Master Resume")
    doc.save(OUTPUT / "conold-chisinahama-resume-master.docx")


if __name__ == "__main__":
    build_concise()
    build_master()
    print("Created concise and master resume DOCX files.")
